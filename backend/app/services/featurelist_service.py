"""FeatureList（功能清单）导出与导入，支持 Markdown 与 Excel 两种格式。

- Markdown：表格形式，多行字段用 <br> 编码换行，单元格内 | 用 \\| 转义，纯文本、Git 友好。
- Excel(xlsx)：所见即所得表格，多行字段自动换行，测试同学在 Excel 里维护更直观。

导出按 format 选择；导入按文件扩展名自动识别。
"""

from io import BytesIO

from app.services.document_parser import DocumentParseError

# 列顺序即导出/导入约定
COLUMNS = [
    ("module", "模块"),
    ("feature", "功能点"),
    ("priority", "优先级"),
    ("description", "功能描述"),
    ("acceptance_criteria", "验收标准"),
    ("constraints", "约束/边界"),
]
HEADER_TO_FIELD = {label: field for field, label in COLUMNS}
COL_WIDTHS = [16, 22, 10, 40, 40, 30]

MEDIA_MD = "text/markdown; charset=utf-8"
MEDIA_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------- Markdown

def _encode_md_cell(value: str) -> str:
    text = (value or "").strip()
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\r\n", "\n").replace("\n", "<br>")


def _decode_md_cell(value: str) -> str:
    text = (value or "").strip()
    text = text.replace("<br>", "\n").replace("<br/>", "\n")
    text = text.replace("\\|", "|").replace("\\\\", "\\")
    return text.strip()


def _split_md_row(line: str) -> list[str]:
    cells: list[str] = []
    buf = ""
    escaped = False
    for ch in line:
        if escaped:
            buf += ch
            escaped = False
            continue
        if ch == "\\":
            buf += ch
            escaped = True
            continue
        if ch == "|":
            cells.append(buf)
            buf = ""
            continue
        buf += ch
    cells.append(buf)
    if cells and cells[0].strip() == "":
        cells = cells[1:]
    if cells and cells[-1].strip() == "":
        cells = cells[:-1]
    return cells


def _is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(set(c.strip()) <= {"-", ":", " "} and "-" in c for c in cells)


def export_featurelist_md(title: str, items: list[dict]) -> str:
    labels = [label for _, label in COLUMNS]
    lines = [
        f"# {title or '功能清单'}",
        "",
        "| " + " | ".join(labels) + " |",
        "| " + " | ".join(["---"] * len(labels)) + " |",
    ]
    for item in items:
        row = [_encode_md_cell(item.get(field, "")) for field, _ in COLUMNS]
        lines.append("| " + " | ".join(row) + " |")
    lines.append("")
    return "\n".join(lines)


def parse_featurelist_md(data: bytes) -> list[dict]:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocumentParseError("无法识别文件编码，请使用 UTF-8")

    table_lines = [ln for ln in text.splitlines() if ln.strip().startswith("|")]
    if not table_lines:
        raise DocumentParseError("未找到 Markdown 表格，请使用「导出清单」得到的格式")

    header_cells = [c.strip() for c in _split_md_row(table_lines[0])]
    col_index = {HEADER_TO_FIELD[name]: i for i, name in enumerate(header_cells) if name in HEADER_TO_FIELD}
    if "feature" not in col_index:
        raise DocumentParseError("FeatureList 缺少「功能点」列")

    items: list[dict] = []
    for line in table_lines[1:]:
        cells = _split_md_row(line)
        if _is_separator_row([c.strip() for c in cells]):
            continue

        def cell(field: str) -> str:
            idx = col_index.get(field)
            if idx is None or idx >= len(cells):
                return ""
            return _decode_md_cell(cells[idx])

        feature = cell("feature")
        if not feature:
            continue
        items.append(_build_item(cell))

    if not items:
        raise DocumentParseError("FeatureList 中未解析到任何功能点")
    return items


# ---------------------------------------------------------------- Excel

def export_featurelist_xlsx(title: str, items: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "FeatureList"

    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)
    for idx, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = COL_WIDTHS[idx - 1]
    ws.freeze_panes = "A2"

    for r, item in enumerate(items, start=2):
        for c, (field, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=item.get(field, "") or "")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def parse_featurelist_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentParseError("FeatureList 解析失败，请确认是 .xlsx 文件") from exc

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise DocumentParseError("FeatureList 为空")

    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col_index = {HEADER_TO_FIELD[name]: i for i, name in enumerate(header) if name in HEADER_TO_FIELD}
    if "feature" not in col_index:
        raise DocumentParseError("FeatureList 缺少「功能点」列")

    items: list[dict] = []
    for row in rows[1:]:
        def cell(field: str) -> str:
            idx = col_index.get(field)
            if idx is None or idx >= len(row):
                return ""
            value = row[idx]
            return str(value).strip() if value is not None else ""

        feature = cell("feature")
        if not feature:
            continue
        items.append(_build_item(cell))

    if not items:
        raise DocumentParseError("FeatureList 中未解析到任何功能点")
    return items


# ---------------------------------------------------------------- shared / dispatch

def _build_item(cell) -> dict:
    priority = cell("priority") or "P1"
    if priority not in ("P0", "P1", "P2"):
        priority = "P1"
    return {
        "module": cell("module"),
        "feature": cell("feature"),
        "description": cell("description"),
        "acceptance_criteria": cell("acceptance_criteria"),
        "constraints": cell("constraints"),
        "priority": priority,
    }


def export_featurelist(title: str, items: list[dict], fmt: str = "xlsx") -> tuple[bytes, str, str]:
    """返回 (内容字节, media_type, 文件扩展名)。"""
    if fmt == "md":
        text = export_featurelist_md(title, items)
        return text.encode("utf-8"), MEDIA_MD, "md"
    content = export_featurelist_xlsx(title, items)
    return content, MEDIA_XLSX, "xlsx"


def parse_featurelist(filename: str, data: bytes) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return parse_featurelist_xlsx(data)
    if name.endswith(".md") or name.endswith(".markdown"):
        return parse_featurelist_md(data)
    raise DocumentParseError("仅支持 .xlsx / .md / .markdown 格式的 FeatureList")
