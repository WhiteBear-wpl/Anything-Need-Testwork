from io import BytesIO

from docx import Document

MAX_FILE_SIZE = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".docx", ".md", ".markdown"}


class DocumentParseError(ValueError):
    pass


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("无法识别 Markdown 文件编码，请使用 UTF-8")


def parse_docx(data: bytes) -> str:
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise DocumentParseError("Word 文件解析失败，请确认是 .docx 格式") from exc

    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    content = "\n\n".join(parts).strip()
    if not content:
        raise DocumentParseError("Word 文件中未提取到文本内容")
    return content


def parse_markdown(data: bytes) -> str:
    content = _decode_text(data).strip()
    if not content:
        raise DocumentParseError("Markdown 文件内容为空")
    return content


def parse_upload(filename: str, data: bytes) -> tuple[str, str]:
    if len(data) > MAX_FILE_SIZE:
        raise DocumentParseError("文件大小不能超过 10MB")

    name = filename.lower()
    if name.endswith(".docx"):
        return parse_docx(data), "docx"
    if name.endswith(".md") or name.endswith(".markdown"):
        return parse_markdown(data), "markdown"

    if name.endswith(".doc"):
        raise DocumentParseError("暂不支持 .doc 格式，请另存为 .docx 后上传")

    raise DocumentParseError("仅支持 .docx、.md、.markdown 格式")


def title_from_filename(filename: str) -> str:
    if "." in filename:
        return filename.rsplit(".", 1)[0]
    return filename
